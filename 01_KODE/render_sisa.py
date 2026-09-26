#!/usr/bin/env python3
"""
Merender citra bertanda yang BELUM punya gambar, lalu menambahkannya ke lembar
tinjauan Excel tanpa menyentuh vonis yang sudah diisi.

Deteksi Penyakit Daun Terung Berbasis Object Detection dengan Model YOLOv11 -

Latar: skrining_label.py cuma merender 40 gambar per pola -- cukup untuk
MENGUKUR persentase, tidak cukup untuk MEMPERBAIKI. Memperbaiki menuntut tiap
citra dilihat satu per satu. Skrip ini merender sisanya.

Yang dirender ditentukan laju ketepatan tiap pola, dari tinjauan 160 sampel
acak kemarin:

    healthy_berlubang  70% tepat  ->  dirender
    tertukar_lesi      60% tepat  ->  dirender
    tertukar_utuh      38% tepat  ->  dirender
    lesi_tampak_utuh    8% tepat  ->  DILEWATI, 120 gambar untuk ~10 temuan

Urutannya menurun menurut skor, karena ketepatan ikut menurun bersama skor.
anotator bisa berhenti kapan saja; yang teratas yang paling berharga.

Pakai:
    python 01_KODE\\render_sisa.py --uji         # self-check
    python 01_KODE\\render_sisa.py               # 384 gambar, ~45 menit
    python 01_KODE\\render_sisa.py --batas 150   # 150 teratas tiap pola
    python 01_KODE\\render_sisa.py --semua-pola  # ikut render lesi_tampak_utuh
"""

import argparse
import csv
import re
import sys
from collections import Counter
from pathlib import Path

import cv2

from audit_prediksi import (NAMA_KELAS, MODEL_TERPILIH, cari_bobot, muat_label,
                            DIR_DATASET, DIR_HASIL)
from skrining_label import nilai, kartu, CONF, POLA

KELUARAN = DIR_HASIL / "skrining_hasil"
XLSX = KELUARAN / "Tinjauan Skrining Label.xlsx"

# Laju ketepatan terukur dari 160 sampel acak, 2026-08-11.
LAJU = {"healthy_berlubang": 0.70, "tertukar_lesi": 0.60,
        "tertukar_utuh": 0.38, "lesi_tampak_utuh": 0.08}
AMBANG_LAYAK = 0.30
VONIS = ["pindah", "biarkan", "buang", "mode_utuh", "campuran"]


def pastikan_openpyxl():
    """
    Diperiksa di AWAL, bukan saat menulis hasil. Kalau ditunda, inferensi 45
    menit sudah terlanjur jalan lalu gagal di baris terakhir.
    """
    try:
        import openpyxl  # noqa: F401
    except ModuleNotFoundError:
        sys.exit("\nopenpyxl belum terpasang. Pasang dulu, lalu ulangi:\n"
                 "    python -m pip install openpyxl\n")


def bersih_deteksi(s):
    return "  +  ".join(f"{a} x{b}" for a, b in re.findall(r"([a-z_]+)x(\d+)", s or ""))


def baca_sudah(xlsx):
    """Kembalikan {nama citra: vonis} dari lembar yang ada. Kosong kalau belum ada."""
    if not xlsx.exists():
        return {}
    from openpyxl import load_workbook
    ws = load_workbook(xlsx)["Tinjauan"]
    return {r[3]: (r[8] or "") for r in ws.iter_rows(min_row=2, values_only=True) if r[3]}


def tulis_xlsx(xlsx, baris):
    """Tulis ulang lembar Tinjauan; vonis lama ikut terbawa di kolom 'vonis'."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    AR = "Arial"
    wb = Workbook()
    ws = wb.active
    ws.title = "Panduan"
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 88
    ws["A1"] = "Lembar Tinjauan Skrining Label"
    ws["A1"].font = Font(AR, 14, bold=True)

    n_isi = sum(1 for b in baris if b["vonis"])
    isi = [("", ""), ("STATUS", "")]
    for p in POLA:
        g = [b for b in baris if b["pola"] == p]
        if g:
            isi.append((p, f"{sum(1 for b in g if b['vonis'])} / {len(g)} selesai"))
    isi += [
        ("TOTAL", f"{n_isi} / {len(baris)} selesai"),
        ("", ""),
        ("CARA PAKAI", ""),
        ("1", "Filter kolom 'vonis' -> centang (Blanks) untuk melihat sisa pekerjaan."),
        ("2", "Kolom 'gambar' menyebut berkas. Buka di 05_HASIL\\skrining_hasil\\<pola>\\"),
        ("3", "Klik sel 'vonis' -> pilih dari dropdown. Jangan diketik manual."),
        ("4", "Simpan Ctrl+S. Tetap .xlsx. JANGAN Save As CSV."),
        ("5", "Baris diurutkan dari skor tertinggi. Yang atas paling mungkin salah label."),
        ("", ""),
        ("ARTI VONIS", ""),
        ("pindah", "kelasnya salah, citra harus pindah ke kelas lain"),
        ("biarkan", "label sudah benar, model yang keliru"),
        ("buang", "ragu, lebih aman disingkirkan dari dataset"),
        ("mode_utuh", "kelas benar TAPI gejala menyebar -> seharusnya 1 kotak seluruh daun"),
        ("campuran", "satu daun memuat DUA kondisi sekaligus (bercak berhalo + lubang bersih)"),
        ("", ""),
        ("PERTANYAAN TIAP POLA", ""),
        ("healthy_berlubang", "Di kotak MERAH ada lubang / bercak nyata? Ada -> pindah."),
        ("tertukar_utuh", "Daunnya BELANG hijau muda-tua (mosaic) atau hijau RATA (healthy)?"),
        ("tertukar_lesi", "Lubang POLOS = insect, lubang BERHALO warna = leaf_spot."),
        ("", "Ada KEDUANYA di satu daun -> campuran."),
        ("lesi_tampak_utuh", "Gejalanya bercak TERPISAH atau warna MENYEBAR rata?"),
        ("", ""),
        ("PENTING", ""),
        ("", "Jangan memindahkan / menghapus berkas citra. Vonis cuma catatan."),
        ("", "Kalau ragu, pilih 'buang' lalu lanjut."),
    ]
    for i, (a, b) in enumerate(isi, start=3):
        ws.cell(i, 1, a).font = Font(AR, 10, bold=(b == "" and a != ""))
        c = ws.cell(i, 2, b)
        c.font = Font(AR, 10)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("Tinjauan")
    KOL = ["no", "pola", "gambar", "citra", "subset", "kelas_folder",
           "skor", "deteksi", "vonis", "catatan"]
    LEB = [22, 20, 44, 40, 9, 22, 8, 32, 14, 28]
    hdr = PatternFill("solid", fgColor="1F3864")
    kuning = PatternFill("solid", fgColor="FFF2CC")
    hijau = PatternFill("solid", fgColor="E2EFDA")
    tipis = Side(style="thin", color="BFBFBF")

    for j, (k, w) in enumerate(zip(KOL, LEB), 1):
        c = ws.cell(1, j, k)
        c.font = Font(AR, 10, bold=True, color="FFFFFF")
        c.fill = hdr
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{len(baris)+1}"

    for i, b in enumerate(baris, 2):
        nilai_baris = [b["no"], b["pola"], b["gambar"], b["citra"], b["subset"],
                       b["kelas_folder"], b["skor"], b["deteksi"], b["vonis"], ""]
        for j, v in enumerate(nilai_baris, 1):
            c = ws.cell(i, j, v)
            c.font = Font(AR, 10)
            c.border = Border(bottom=tipis)
            if j == 7:
                c.number_format = "0.00"
            if j in (9, 10):
                c.fill = hijau if b["vonis"] else kuning
        if b["subset"] == "test":
            ws.cell(i, 5).font = Font(AR, 10, bold=True, color="C00000")

    dv = DataValidation(type="list", formula1='"' + ",".join(VONIS) + '"',
                        allow_blank=True, showDropDown=False)
    dv.errorTitle = "Vonis tidak dikenali"
    dv.error = "Pilih salah satu dari daftar."
    ws.add_data_validation(dv)
    dv.add(f"I2:I{len(baris)+1}")
    wb.save(xlsx)


def jalankan(args):
    pastikan_openpyxl()                 # gagal sekarang, bukan setelah 45 menit
    from ultralytics import YOLO

    antrian = KELUARAN / "antrian.csv"
    if not antrian.exists():
        sys.exit(f"Tidak ada {antrian}. Jalankan skrining_label.py dulu.")
    semua = list(csv.DictReader(open(antrian, encoding="utf-8")))
    sudah = baca_sudah(XLSX)
    print(f"bertanda total : {len(semua)}")
    print(f"sudah ada di lembar : {len(sudah)}")

    layak = set(POLA) if args.semua_pola else {p for p in POLA if LAJU[p] >= AMBANG_LAYAK}
    dilewati = set(POLA) - layak
    if dilewati:
        print(f"pola dilewati  : {', '.join(sorted(dilewati))} (laju ketepatan rendah)")

    sisa = [r for r in semua if r["citra"] not in sudah and r["pola"] in layak]
    sisa.sort(key=lambda r: (r["pola"], -float(r["skor"])))
    if args.batas:
        per = {}
        pilih = []
        for r in sisa:
            per[r["pola"]] = per.get(r["pola"], 0) + 1
            if per[r["pola"]] <= args.batas:
                pilih.append(r)
        sisa = pilih
    print(f"akan dirender  : {len(sisa)}")
    for p, n in sorted(Counter(r["pola"] for r in sisa).items()):
        print(f"   {p:20s} {n}")
    if not sisa:
        print("\nTidak ada yang perlu dirender.")
        return

    split = args.split
    bobot = cari_bobot(split, args.model or MODEL_TERPILIH[split])
    if bobot is None:
        sys.exit("Bobot tidak ketemu.")
    print(f"\nbobot : {bobot.name}\nimgsz : {args.imgsz}\n")
    model = YOLO(str(bobot))

    akar = DIR_DATASET / f"dataset_{split}"
    peta = {}
    for sub in ["train", "val", "test"]:
        d = akar / "images" / sub
        if d.is_dir():
            for p in d.iterdir():
                peta[p.name] = (p, akar / "labels" / sub / f"{p.stem}.txt")

    # nomor lanjut dari yang sudah ada di tiap folder pola
    mulai = {}
    for p in POLA:
        d = KELUARAN / p
        ada = [int(f.stem.split("_", 1)[0]) for f in d.glob("*.jpg")
               if f.stem.split("_", 1)[0].isdigit()] if d.is_dir() else []
        mulai[p] = max(ada) if ada else 0

    baru, gagal = [], []
    for n, r in enumerate(sisa, 1):
        if r["citra"] not in peta:
            gagal.append(r["citra"]); continue
        pimg, plbl = peta[r["citra"]]
        citra = cv2.imread(str(pimg), cv2.IMREAD_COLOR)
        if citra is None:
            gagal.append(r["citra"]); continue
        h, w = citra.shape[:2]
        label = muat_label(plbl, w, h)
        kf = Counter(b[0] for b in label).most_common(1)[0][0] if label else -1

        res = model.predict(str(pimg), conf=CONF, imgsz=args.imgsz, verbose=False)[0]
        pred = [(int(c), *map(float, xy), float(cf))
                for c, xy, cf in zip(res.boxes.cls.tolist(), res.boxes.xyxy.tolist(),
                                     res.boxes.conf.tolist())]
        skor, pola, _ = nilai(kf, pred, float(w * h))
        pola = pola or r["pola"]

        mulai[pola] += 1
        urut = mulai[pola]
        tuj = KELUARAN / pola
        tuj.mkdir(parents=True, exist_ok=True)
        nama = f"{urut:03d}_{Path(r['citra']).stem}.jpg"
        cv2.imwrite(str(tuj / nama),
                    kartu(citra, label, pred, kf, r["citra"], pola, skor),
                    [cv2.IMWRITE_JPEG_QUALITY, 88])

        baru.append({"no": f"{pola}/{urut:03d}", "pola": pola, "gambar": nama,
                     "citra": r["citra"], "subset": r["subset"],
                     "kelas_folder": r["kelas_folder"], "skor": float(r["skor"]),
                     "deteksi": bersih_deteksi(r["deteksi_asing"]), "vonis": ""})
        if n % 25 == 0:
            print(f"  {n}/{len(sisa)}")

    if gagal:
        print(f"\n{len(gagal)} citra gagal dibaca: {gagal[:5]}")

    # gabung: baris lama (dengan vonisnya) + baris baru
    lama = []
    if XLSX.exists():
        from openpyxl import load_workbook
        ws = load_workbook(XLSX)["Tinjauan"]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[3]:
                lama.append({"no": r[0], "pola": r[1], "gambar": r[2], "citra": r[3],
                             "subset": r[4], "kelas_folder": r[5], "skor": r[6],
                             "deteksi": r[7], "vonis": r[8] or ""})
    gabung = lama + baru
    urut_pola = {p: i for i, p in enumerate(POLA)}
    gabung.sort(key=lambda b: (urut_pola.get(b["pola"], 9), -float(b["skor"] or 0)))
    tulis_xlsx(XLSX, gabung)

    print(f"\n{len(baru)} gambar baru ditulis ke {KELUARAN}")
    print(f"lembar diperbarui: {len(gabung)} baris, "
          f"{sum(1 for b in gabung if b['vonis'])} sudah bervonis, "
          f"{sum(1 for b in gabung if not b['vonis'])} menunggu")
    print("\nVonis lama TIDAK tersentuh. Buka lembarnya, filter kolom 'vonis'")
    print("-> centang (Blanks), lalu kerjakan dari atas.")


def uji():
    lolos = True

    def cek(n, s):
        nonlocal lolos
        print(f"  {'ok  ' if s else 'GAGAL'}  {n}")
        lolos = lolos and s

    print("bersih_deteksi()")
    cek("ubah format", bersih_deteksi("leaf_spot_diseasex6; healthy_leafx1")
        == "leaf_spot_disease x6  +  healthy_leaf x1")
    cek("kosong aman", bersih_deteksi("") == "")
    cek("None aman", bersih_deteksi(None) == "")

    print("pemilihan pola layak")
    layak = {p for p in POLA if LAJU[p] >= AMBANG_LAYAK}
    cek("lesi_tampak_utuh dilewati", "lesi_tampak_utuh" not in layak)
    cek("tiga pola lain ikut", len(layak) == 3)
    cek("healthy_berlubang ikut", "healthy_berlubang" in layak)

    print("tulis_xlsx() + baca_sudah()")
    try:
        import openpyxl  # noqa: F401
    except ModuleNotFoundError:
        print("  DILEWATI  openpyxl belum terpasang -> python -m pip install openpyxl")
        print("\n" + ("SEMUA UJI LOLOS (kecuali bagian openpyxl)" if lolos else "ADA YANG GAGAL"))
        return 0 if lolos else 1
    import tempfile
    d = Path(tempfile.mkdtemp())
    f = d / "uji.xlsx"
    baris = [{"no": "healthy_berlubang/001", "pola": "healthy_berlubang",
              "gambar": "001_a.jpg", "citra": "a.jpg", "subset": "test",
              "kelas_folder": "healthy_leaf", "skor": 2.5,
              "deteksi": "insect_pest_disease x3", "vonis": "pindah"},
             {"no": "tertukar_lesi/001", "pola": "tertukar_lesi",
              "gambar": "001_b.jpg", "citra": "b.jpg", "subset": "train",
              "kelas_folder": "leaf_spot_disease", "skor": 1.0,
              "deteksi": "insect_pest_disease x2", "vonis": ""}]
    tulis_xlsx(f, baris)
    cek("berkas tertulis", f.exists())
    kembali = baca_sudah(f)
    cek("dua baris terbaca", len(kembali) == 2)
    cek("vonis terjaga", kembali["a.jpg"] == "pindah")
    cek("vonis kosong tetap kosong", kembali["b.jpg"] == "")

    from openpyxl import load_workbook
    wb = load_workbook(f)
    cek("dua sheet", wb.sheetnames == ["Panduan", "Tinjauan"])
    ws = wb["Tinjauan"]
    cek("validasi dropdown ada", len(ws.data_validations.dataValidation) == 1)
    cek("freeze panes", ws.freeze_panes == "A2")

    print("\n" + ("SEMUA UJI LOLOS" if lolos else "ADA YANG GAGAL"))
    return 0 if lolos else 1


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--split", default="80_10_10", choices=["80_10_10", "70_15_15"])
    p.add_argument("--model", default=None)
    p.add_argument("--imgsz", type=int, default=640,
                   help="harus sama dengan imgsz saat model dilatih")
    p.add_argument("--batas", type=int, default=0, help="maksimum gambar per pola")
    p.add_argument("--semua-pola", action="store_true",
                   help="ikut render lesi_tampak_utuh (laju ketepatan cuma 8%%)")
    p.add_argument("--uji", action="store_true")
    a = p.parse_args()
    if a.uji:
        sys.exit(uji())
    jalankan(a)


if __name__ == "__main__":
    main()
