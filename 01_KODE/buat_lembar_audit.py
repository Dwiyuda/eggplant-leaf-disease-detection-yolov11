#!/usr/bin/env python3
"""
Mengubah prioritas.csv jadi lembar Excel siap isi, dengan dropdown vonis.

Deteksi Penyakit Daun Terung Berbasis Object Detection dengan Model YOLOv11 -

Formatnya mengikuti lembar tinjauan skrining yang sudah terbukti nyaman dipakai:
header beku, filter, kolom isian berwarna, dan dropdown supaya vonis tidak bisa
salah ketik.

Satu tambahan penting: baris berkelompok `acak` diberi warna berbeda. Cuma
kelompok itu yang boleh dipakai menghitung persentase -- kelompok `terburuk`
sengaja bias karena dipilih sebagai kasus paling parah.

Aman dijalankan berulang. Vonis yang sudah diisi TIDAK akan hilang.

Pakai:
    python 01_KODE\\buat_lembar_audit.py --uji
    python 01_KODE\\buat_lembar_audit.py
"""

import argparse
import csv
import sys
from collections import Counter
from pathlib import Path

from audit_prediksi import DIR_DATASET

AKAR = DIR_DATASET.parent
DIR_AUDIT = AKAR / "06_AUDIT_KOTAK"
CSV = DIR_AUDIT / "prioritas.csv"
XLSX = DIR_AUDIT / "Audit Kotak.xlsx"

VONIS = ["banyak_palsu", "sebagian_palsu", "sudah_benar", "kurang_kotak"]
ARTI = {
    "banyak_palsu": "sebagian besar kotak lama TIDAK ada bercak/lubangnya",
    "sebagian_palsu": "campur -- ada yang benar, ada yang mengada-ada",
    "sudah_benar": "kotaknya wajar, paling cuma perlu digeser sedikit",
    "kurang_kotak": "bercaknya justru LEBIH banyak dari yang dilabeli",
}

KOL = ["no", "kelas", "kelompok", "gambar", "subset", "n_label",
       "n_prediksi", "tidak_ditemukan", "persen_terlewat", "vonis_kotak", "catatan"]
LEBAR = [6, 22, 10, 42, 8, 8, 10, 14, 14, 16, 34]


def baca_vonis_lama(xlsx):
    """{nama citra: (vonis, catatan)} dari lembar yang sudah ada."""
    if not xlsx.exists():
        return {}
    from openpyxl import load_workbook
    ws = load_workbook(xlsx)["Audit"]
    hdr = [c.value for c in ws[1]]
    try:
        i_g, i_v, i_c = hdr.index("gambar"), hdr.index("vonis_kotak"), hdr.index("catatan")
    except ValueError:
        return {}
    return {r[i_g]: (r[i_v] or "", r[i_c] or "")
            for r in ws.iter_rows(min_row=2, values_only=True) if r[i_g]}


def susun(rows, lama):
    """Urut: kelompok acak dulu (itu yang memberi angka), lalu kelas, lalu skor."""
    out = []
    for r in rows:
        v, c = lama.get(r["citra"], ("", ""))
        out.append({**r, "vonis_kotak": r.get("vonis_kotak") or v,
                    "catatan": r.get("catatan") or c})
    out.sort(key=lambda r: (r["kelompok"] != "acak", r["kelas"],
                            -int(r["tidak_ditemukan"])))
    for i, r in enumerate(out, 1):
        r["no"] = i
        r["gambar"] = r["citra"]
    return out


def tulis(rows, xlsx):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    AR = "Arial"
    wb = Workbook()

    ws = wb.active
    ws.title = "Panduan"
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 92
    ws["A1"] = "Audit Kotak — leaf_spot & insect"
    ws["A1"].font = Font(AR, 14, bold=True)
    ws["A2"] = "penelitian"
    ws["A2"].font = Font(AR, 10, italic=True)

    n_acak = sum(1 for r in rows if r["kelompok"] == "acak")
    isi = [("", ""), ("PERTANYAANNYA", ""),
           ("", "Recall leaf_spot cuma 0,2297 — model menemukan kurang dari seperempat"),
           ("", "bercak yang dilabeli. Apakah bercaknya memang ada tapi terlalu kecil,"),
           ("", "atau sebagian kotak itu dibuat aturan warna padahal tidak ada apa-apa?"),
           ("", ""), ("CARA KERJA", ""),
           ("1", "Buka citra di Roboflow, koreksi kotaknya."),
           ("2", "Cari barisnya di sheet 'Audit' — Ctrl+F, tempel nama berkasnya."),
           ("3", "Klik sel kolom 'vonis_kotak' → pilih dari dropdown."),
           ("4", "Kolom 'catatan' bebas, boleh dikosongkan."),
           ("5", "Simpan Ctrl+S. Tetap .xlsx."),
           ("", ""), ("ARTI VONIS", "")]
    isi += [(v, ARTI[v]) for v in VONIS]
    isi += [("", ""), ("KELOMPOK — INI PENTING", ""),
            ("acak", f"{n_acak} citra, diambil acak. HANYA kelompok ini yang boleh"),
            ("", "dipakai menghitung persentase. Barisnya berlatar hijau."),
            ("terburuk", f"{len(rows)-n_acak} citra, sengaja dipilih sebagai kasus paling parah."),
            ("", "Gunanya membuktikan gejalanya nyata, BUKAN mengukur seberapa sering."),
            ("", "Kalau dipakai menghitung, angkanya pasti terlalu tinggi."),
            ("", ""), ("KALAU WAKTU MEPET", ""),
            ("", "Kerjakan baris berlatar hijau (acak) lebih dulu — itu yang memberi angka."),
            ("", "Kelompok terburuk boleh menyusul, atau dilewati sama sekali."),
            ("", ""), ("MEMBACA KOLOM", ""),
            ("n_label", "berapa kotak yang dibuat autolabel.py pada citra itu"),
            ("n_prediksi", "berapa kotak yang ditemukan model"),
            ("tidak_ditemukan", "kotak label yang tidak ditemukan model — makin besar makin curiga"),
            ("persen_terlewat", "tidak_ditemukan dibagi n_label")]
    for i, (a, b) in enumerate(isi, start=4):
        ws.cell(i, 1, a).font = Font(AR, 10, bold=(b == "" and a != ""))
        c = ws.cell(i, 2, b)
        c.font = Font(AR, 10)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("Audit")
    hdr = PatternFill("solid", fgColor="1F3864")
    kuning = PatternFill("solid", fgColor="FFF2CC")
    hijau_isi = PatternFill("solid", fgColor="E2EFDA")
    hijau_acak = PatternFill("solid", fgColor="F2F8F0")
    tipis = Side(style="thin", color="BFBFBF")

    for j, (k, w) in enumerate(zip(KOL, LEBAR), 1):
        c = ws.cell(1, j, k)
        c.font = Font(AR, 10, bold=True, color="FFFFFF")
        c.fill = hdr
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(KOL))}{len(rows)+1}"

    for i, r in enumerate(rows, 2):
        acak = r["kelompok"] == "acak"
        for j, k in enumerate(KOL, 1):
            v = r[k]
            if k in ("n_label", "n_prediksi", "tidak_ditemukan"):
                v = int(v)
            elif k == "persen_terlewat":
                v = float(v)
            c = ws.cell(i, j, v)
            c.font = Font(AR, 10)
            c.border = Border(bottom=tipis)
            if k == "persen_terlewat":
                c.number_format = "0%"
            if k in ("vonis_kotak", "catatan"):
                c.fill = hijau_isi if r["vonis_kotak"] else kuning
            elif acak:
                c.fill = hijau_acak
        if acak:
            ws.cell(i, 3).font = Font(AR, 10, bold=True, color="1E7145")
        if r["subset"] == "test":
            ws.cell(i, 5).font = Font(AR, 10, bold=True, color="C00000")

    dv = DataValidation(type="list", formula1='"' + ",".join(VONIS) + '"',
                        allow_blank=True, showDropDown=False)
    dv.errorTitle = "Vonis tidak dikenali"
    dv.error = "Pilih salah satu dari daftar."
    ws.add_data_validation(dv)
    kol_v = get_column_letter(KOL.index("vonis_kotak") + 1)
    dv.add(f"{kol_v}2:{kol_v}{len(rows)+1}")
    wb.save(xlsx)


def uji():
    lolos = True

    def cek(n, s):
        nonlocal lolos
        print(f"  {'ok  ' if s else 'GAGAL'}  {n}")
        lolos = lolos and s

    print("susun()")
    rows = [
        {"citra": "a.jpg", "kelas": "leaf_spot_disease", "kelompok": "terburuk",
         "subset": "train", "n_label": "20", "n_prediksi": "2",
         "tidak_ditemukan": "18", "persen_terlewat": "0.9", "vonis_kotak": "", "catatan": ""},
        {"citra": "b.jpg", "kelas": "leaf_spot_disease", "kelompok": "acak",
         "subset": "val", "n_label": "10", "n_prediksi": "8",
         "tidak_ditemukan": "2", "persen_terlewat": "0.2", "vonis_kotak": "", "catatan": ""},
        {"citra": "c.jpg", "kelas": "insect_pest_disease", "kelompok": "acak",
         "subset": "test", "n_label": "9", "n_prediksi": "1",
         "tidak_ditemukan": "8", "persen_terlewat": "0.889", "vonis_kotak": "", "catatan": ""},
    ]
    out = susun(rows, {})
    cek("kelompok acak ditaruh di atas",
        out[0]["kelompok"] == "acak" and out[-1]["kelompok"] == "terburuk")
    cek("nomor urut dibuat ulang", [r["no"] for r in out] == [1, 2, 3])
    cek("kolom gambar mengikuti nama citra", out[0]["gambar"] == out[0]["citra"])
    cek("semua baris terbawa", len(out) == 3)

    print("vonis lama dipertahankan")
    out = susun(rows, {"a.jpg": ("banyak_palsu", "kotak di urat daun")})
    a = [r for r in out if r["citra"] == "a.jpg"][0]
    cek("vonis lama terbawa", a["vonis_kotak"] == "banyak_palsu")
    cek("catatan lama terbawa", a["catatan"] == "kotak di urat daun")
    cek("baris lain tetap kosong",
        all(r["vonis_kotak"] == "" for r in out if r["citra"] != "a.jpg"))

    print("tulis() + baca ulang")
    import tempfile
    f = Path(tempfile.mkdtemp()) / "uji.xlsx"
    tulis(susun(rows, {"a.jpg": ("sudah_benar", "oke")}), f)
    cek("berkas tertulis", f.exists())
    kembali = baca_vonis_lama(f)
    cek("tiga baris terbaca", len(kembali) == 3)
    cek("vonis bertahan lewat tulis-baca", kembali["a.jpg"][0] == "sudah_benar")
    from openpyxl import load_workbook
    wb = load_workbook(f)
    cek("dua sheet", wb.sheetnames == ["Panduan", "Audit"])
    ws = wb["Audit"]
    cek("dropdown terpasang", len(ws.data_validations.dataValidation) == 1)
    cek("header beku di kolom E", ws.freeze_panes == "E2")

    print("\n" + ("SEMUA UJI LOLOS" if lolos else "ADA YANG GAGAL"))
    return 0 if lolos else 1


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--uji", action="store_true")
    a = p.parse_args()
    if a.uji:
        sys.exit(uji())

    try:
        import openpyxl  # noqa: F401
    except ModuleNotFoundError:
        sys.exit("\nopenpyxl belum terpasang:  python -m pip install openpyxl\n")
    if not CSV.exists():
        sys.exit(f"Tidak ada {CSV}. Jalankan siapkan_audit_kotak.py dulu.")

    rows = list(csv.DictReader(open(CSV, encoding="utf-8-sig")))
    lama = baca_vonis_lama(XLSX)
    if lama:
        print(f"vonis yang sudah ada : {sum(1 for v in lama.values() if v[0])} — dipertahankan")
    out = susun(rows, lama)
    tulis(out, XLSX)

    c = Counter(r["kelompok"] for r in out)
    print(f"\n{XLSX.name} dibuat — {len(out)} baris")
    print(f"  acak     {c['acak']:3d}  (berlatar hijau, ini yang dipakai menghitung)")
    print(f"  terburuk {c['terburuk']:3d}  (kasus paling parah, untuk membuktikan gejalanya)")
    print(f"  sudah bervonis {sum(1 for r in out if r['vonis_kotak'])}")
    print("\nBuka sheet 'Audit'. Cari baris pakai Ctrl+F, tempel nama berkas dari Roboflow.")


if __name__ == "__main__":
    main()
