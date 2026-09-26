#!/usr/bin/env python3
"""
Memisahkan gambar tinjauan jadi dua folder: yang BELUM divonis dan yang SUDAH,
lalu menomori ulang yang belum agar urutan berkas sama persis dengan urutan
baris di Excel.

Deteksi Penyakit Daun Terung Berbasis Object Detection dengan Model YOLOv11 -

MASALAH YANG DIPECAHKAN. Setelah render_sisa.py, satu folder memuat campuran
yang sudah dan belum divonis, dan penomorannya tidak berurutan dengan baris
Excel -- baris pertama yang menunggu justru berkas nomor 041. Meninjau jadi
lompat-lompat dan gampang tertukar.

SESUDAH SKRIP INI:

    skrining_hasil/
    ├── BELUM_DIVONIS/<pola>/   001, 002, 003, ...  urut skor, urut baris Excel
    └── SUDAH_DIVONIS/<pola>/   arsip, nama asli dipertahankan

Buka satu folder, urutkan menurut nama, kerjakan dari atas. Baris ke-1 pada
Excel yang difilter = berkas 001, baris ke-2 = berkas 002, dan seterusnya.

AMAN DIJALANKAN BERULANG. Berkas dicari di semua lokasi yang mungkin, jadi
setelah anotator memvonis sebagian, jalankan lagi dan yang baru bervonis otomatis
pindah ke SUDAH_DIVONIS sementara sisanya dinomori ulang.

Pakai:
    python 01_KODE\\rapikan_tinjauan.py --uji        # self-check
    python 01_KODE\\rapikan_tinjauan.py              # KERING: laporan saja
    python 01_KODE\\rapikan_tinjauan.py --terapkan   # sungguhan
"""

import argparse
import shutil
import sys
from collections import Counter
from pathlib import Path

from audit_prediksi import DIR_HASIL

KELUARAN = DIR_HASIL / "skrining_hasil"
XLSX = KELUARAN / "Tinjauan Skrining Label.xlsx"
BELUM, SUDAH = "BELUM_DIVONIS", "SUDAH_DIVONIS"
POLA = ["healthy_berlubang", "lesi_tampak_utuh", "tertukar_utuh", "tertukar_lesi"]
VONIS = ["pindah", "biarkan", "buang", "mode_utuh", "campuran"]


def cari_berkas(nama, pola):
    """Cari gambar di semua lokasi yang mungkin. Mendukung jalan berulang."""
    for kandidat in (KELUARAN / BELUM / pola / nama,
                     KELUARAN / SUDAH / pola / nama,
                     KELUARAN / pola / nama):
        if kandidat.exists():
            return kandidat
    return None


def nama_baru(urut, nama_lama):
    """Ganti awalan nomor lama dengan nomor urut baru."""
    sisa = nama_lama.split("_", 1)[1] if "_" in nama_lama else nama_lama
    return f"{urut:03d}_{sisa}"


def susun(rows):
    """
    Kembalikan rencana pemindahan.
    Yang belum divonis diurutkan skor menurun per pola lalu dinomori 001..N,
    supaya nomor berkas mengikuti urutan baris Excel.
    """
    rencana = []
    for pola in POLA:
        belum = sorted([r for r in rows if r["pola"] == pola and not r["vonis"]],
                       key=lambda r: -(r["skor"] or 0))
        for i, r in enumerate(belum, 1):
            rencana.append({**r, "tujuan_dir": BELUM, "nama_baru": nama_baru(i, r["gambar"]),
                            "no_baru": f"{pola}/{i:03d}"})
        for r in [r for r in rows if r["pola"] == pola and r["vonis"]]:
            rencana.append({**r, "tujuan_dir": SUDAH, "nama_baru": r["gambar"],
                            "no_baru": r["no"]})
    return rencana


def baca_rows():
    from openpyxl import load_workbook
    if not XLSX.exists():
        sys.exit(f"Tidak ada {XLSX}")
    ws = load_workbook(XLSX)["Tinjauan"]
    return [{"no": r[0], "pola": r[1], "gambar": r[2], "citra": r[3], "subset": r[4],
             "kelas_folder": r[5], "skor": r[6], "deteksi": r[7],
             "vonis": (r[8] or "").strip(), "catatan": r[9] or ""}
            for r in ws.iter_rows(min_row=2, values_only=True) if r[3]]


def tulis_xlsx(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    AR = "Arial"
    wb = Workbook()
    ws = wb.active
    ws.title = "Panduan"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 86
    ws["A1"] = "Lembar Tinjauan Skrining Label"
    ws["A1"].font = Font(AR, 14, bold=True)

    belum = [r for r in rows if not r["vonis"]]
    isi = [("", ""), ("DI MANA GAMBARNYA", ""),
           ("", "Yang BELUM divonis -> 05_HASIL\\skrining_hasil\\BELUM_DIVONIS\\<pola>\\"),
           ("", "Yang SUDAH divonis -> 05_HASIL\\skrining_hasil\\SUDAH_DIVONIS\\<pola>\\"),
           ("", ""),
           ("", "Nomor berkas di BELUM_DIVONIS URUT dengan baris di sini."),
           ("", "Baris pertama yang menunggu = berkas 001, baris kedua = 002, dst."),
           ("", ""), ("SISA PEKERJAAN", "")]
    for p in POLA:
        n = sum(1 for r in belum if r["pola"] == p)
        if n:
            isi.append((p, f"{n} gambar menunggu"))
    isi += [("TOTAL", f"{len(belum)} menunggu dari {len(rows)}"),
            ("", ""), ("CARA PAKAI", ""),
            ("1", "Filter kolom 'vonis' -> centang (Blanks) saja."),
            ("2", "Buka folder BELUM_DIVONIS\\<pola>\\, urutkan menurut Nama."),
            ("3", "Kerjakan dari 001 turun. Urutannya sama dengan baris di sini."),
            ("4", "Klik sel 'vonis' -> pilih dari dropdown. Jangan diketik manual."),
            ("5", "Simpan Ctrl+S. Tetap .xlsx, jangan Save As CSV."),
            ("6", "Boleh berhenti kapan saja. Yang atas paling berharga."),
            ("", ""), ("ARTI VONIS", ""),
            ("pindah", "kelasnya salah, citra harus pindah ke kelas lain"),
            ("biarkan", "label sudah benar, model yang keliru"),
            ("buang", "ragu, lebih aman disingkirkan dari dataset"),
            ("mode_utuh", "kelas benar TAPI gejala menyebar -> seharusnya 1 kotak seluruh daun"),
            ("campuran", "satu daun memuat DUA kondisi (bercak berhalo + lubang bersih)"),
            ("", ""), ("PERTANYAAN TIAP POLA", ""),
            ("healthy_berlubang", "Di kotak MERAH ada lubang / bercak nyata? Ada -> pindah."),
            ("tertukar_utuh", "Daunnya BELANG hijau muda-tua (mosaic) atau hijau RATA (healthy)?"),
            ("tertukar_lesi", "Lubang POLOS = insect, BERHALO warna = leaf_spot."),
            ("", "Ada KEDUANYA di satu daun -> campuran."),
            ("", ""), ("PENTING", ""),
            ("", "Jangan memindahkan / menghapus berkas citra. Vonis cuma catatan."),
            ("", "Setelah memvonis sebagian, boleh jalankan rapikan_tinjauan.py lagi"),
            ("", "supaya yang sudah selesai pindah ke SUDAH_DIVONIS dan sisanya dinomori ulang.")]
    for i, (a, b) in enumerate(isi, start=3):
        ws.cell(i, 1, a).font = Font(AR, 10, bold=(b == "" and a != ""))
        c = ws.cell(i, 2, b)
        c.font = Font(AR, 10)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("Tinjauan")
    KOL = ["no", "pola", "gambar", "citra", "subset", "kelas_folder",
           "skor", "deteksi", "vonis", "catatan"]
    LEB = [22, 20, 44, 40, 9, 22, 8, 32, 14, 28]
    hdr = PatternFill("solid", fgColor="1F3864")
    kuning = PatternFill("solid", fgColor="FFF2CC")
    hijau = PatternFill("solid", fgColor="E2EFDA")
    tipis = Side(style="thin", color="BFBFBF")
    for j, (k, w) in enumerate(zip(KOL, LEB), 1):
        c = ws.cell(1, j, k)
        c.font = Font(AR, 10, bold=True, color="FFFFFF")
        c.fill = hdr
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{len(rows)+1}"
    for i, r in enumerate(rows, 2):
        for j, v in enumerate([r["no"], r["pola"], r["gambar"], r["citra"], r["subset"],
                               r["kelas_folder"], r["skor"], r["deteksi"],
                               r["vonis"], r["catatan"]], 1):
            c = ws.cell(i, j, v)
            c.font = Font(AR, 10)
            c.border = Border(bottom=tipis)
            if j == 7:
                c.number_format = "0.00"
            if j in (9, 10):
                c.fill = hijau if r["vonis"] else kuning
        if r["subset"] == "test":
            ws.cell(i, 5).font = Font(AR, 10, bold=True, color="C00000")
    dv = DataValidation(type="list", formula1='"' + ",".join(VONIS) + '"',
                        allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(f"I2:I{len(rows)+1}")
    wb.save(XLSX)


def uji():
    lolos = True

    def cek(n, s):
        nonlocal lolos
        print(f"  {'ok  ' if s else 'GAGAL'}  {n}")
        lolos = lolos and s

    print("nama_baru()")
    cek("ganti awalan nomor", nama_baru(7, "041_Mosaic Virus Disease Image 2.jpg")
        == "007_Mosaic Virus Disease Image 2.jpg")
    cek("tanpa awalan nomor", nama_baru(3, "Gambar.jpg") == "003_Gambar.jpg")
    cek("nomor tiga digit", nama_baru(210, "001_a.jpg") == "210_a.jpg")
    cek("nama bergaris bawah ganda",
        nama_baru(1, "099_Leaf_Spot_5.jpg") == "001_Leaf_Spot_5.jpg")

    print("susun()")
    rows = [
        {"no": "healthy_berlubang/041", "pola": "healthy_berlubang", "gambar": "041_a.jpg",
         "citra": "a.jpg", "skor": 1.0, "vonis": ""},
        {"no": "healthy_berlubang/042", "pola": "healthy_berlubang", "gambar": "042_b.jpg",
         "citra": "b.jpg", "skor": 5.0, "vonis": ""},
        {"no": "healthy_berlubang/001", "pola": "healthy_berlubang", "gambar": "001_c.jpg",
         "citra": "c.jpg", "skor": 3.0, "vonis": "pindah"},
        {"no": "tertukar_lesi/041", "pola": "tertukar_lesi", "gambar": "041_d.jpg",
         "citra": "d.jpg", "skor": 2.0, "vonis": ""},
    ]
    r = susun(rows)
    belum = [x for x in r if x["tujuan_dir"] == BELUM]
    sudah = [x for x in r if x["tujuan_dir"] == SUDAH]
    cek("tiga belum, satu sudah", len(belum) == 3 and len(sudah) == 1)
    hb = [x for x in belum if x["pola"] == "healthy_berlubang"]
    cek("skor tertinggi jadi nomor 001",
        hb[0]["citra"] == "b.jpg" and hb[0]["nama_baru"].startswith("001_"))
    cek("skor lebih rendah jadi 002",
        hb[1]["citra"] == "a.jpg" and hb[1]["nama_baru"].startswith("002_"))
    cek("penomoran ulang per pola",
        [x for x in belum if x["pola"] == "tertukar_lesi"][0]["nama_baru"].startswith("001_"))
    cek("yang sudah divonis namanya tetap", sudah[0]["nama_baru"] == "001_c.jpg")
    cek("no_baru ikut diperbarui", hb[0]["no_baru"] == "healthy_berlubang/001")
    cek("semua baris terbawa", len(r) == len(rows))

    print("\n" + ("SEMUA UJI LOLOS" if lolos else "ADA YANG GAGAL"))
    return 0 if lolos else 1


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--terapkan", action="store_true", help="sungguhan memindahkan berkas")
    p.add_argument("--uji", action="store_true")
    a = p.parse_args()
    if a.uji:
        sys.exit(uji())

    try:
        import openpyxl  # noqa: F401
    except ModuleNotFoundError:
        sys.exit("\nopenpyxl belum terpasang:  python -m pip install openpyxl\n")

    rows = baca_rows()
    rencana = susun(rows)
    belum = [r for r in rencana if r["tujuan_dir"] == BELUM]
    sudah = [r for r in rencana if r["tujuan_dir"] == SUDAH]

    print(f"total baris : {len(rows)}")
    print(f"  BELUM_DIVONIS : {len(belum)}")
    print(f"  SUDAH_DIVONIS : {len(sudah)}\n")
    print(f"{'pola':22s} {'belum':>6s} {'sudah':>6s}")
    for pl in POLA:
        b = sum(1 for r in belum if r["pola"] == pl)
        s = sum(1 for r in sudah if r["pola"] == pl)
        if b or s:
            print(f"{pl:22s} {b:6d} {s:6d}")

    hilang = [r for r in rencana if cari_berkas(r["gambar"], r["pola"]) is None]
    if hilang:
        print(f"\n{len(hilang)} berkas tidak ketemu, contoh: "
              f"{[r['gambar'] for r in hilang[:3]]}")

    if not a.terapkan:
        print("\n" + "-" * 74)
        print("MODE KERING. Belum ada berkas yang dipindah.")
        print("Kalau rencana di atas benar, jalankan:")
        print("    python 01_KODE\\rapikan_tinjauan.py --terapkan")
        print("-" * 74)
        return

    pindah = 0
    for r in rencana:
        asal = cari_berkas(r["gambar"], r["pola"])
        if asal is None:
            continue
        tuj = KELUARAN / r["tujuan_dir"] / r["pola"] / r["nama_baru"]
        if asal.resolve() == tuj.resolve():
            continue
        tuj.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(asal), str(tuj))
        pindah += 1

    baru = []
    for r in rencana:
        baru.append({**r, "gambar": r["nama_baru"], "no": r["no_baru"]})
    urut = {p: i for i, p in enumerate(POLA)}
    baru.sort(key=lambda r: (0 if not r["vonis"] else 1, urut.get(r["pola"], 9),
                             -(r["skor"] or 0)))
    tulis_xlsx(baru)

    for pl in POLA:                       # bersihkan folder lama yang kosong
        d = KELUARAN / pl
        if d.is_dir() and not any(d.iterdir()):
            d.rmdir()

    print(f"\n{pindah} berkas dipindah.")
    print(f"lembar diperbarui: {len(baru)} baris\n")
    print("Sekarang:")
    print(f"  BELUM_DIVONIS\\<pola>\\  -> {len(belum)} gambar, dinomori 001 dst urut skor")
    print(f"  SUDAH_DIVONIS\\<pola>\\  -> {len(sudah)} gambar, arsip")
    print("\nBuka Excel, filter kolom 'vonis' -> (Blanks). Baris pertama = berkas 001.")


if __name__ == "__main__":
    main()
